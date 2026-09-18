"""Real PPTX API evaluation. Missing E2E capabilities are failures, not mock passes."""
import argparse
import hashlib
import json
import sys
import tempfile
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from pptx import Presentation
from fastapi.testclient import TestClient
from adaptive_learning import api
from adaptive_learning.storage import StorageManager
from main import app

# Reviewer annotations; prerequisite edges express curriculum expectations.
ANNOTATIONS = [
 (3,'Trí tuệ nhân tạo','artificial intelligence',[]),
 (4,'AI phân tích','dữ liệu trong quá khứ',[]),
 (5,'Định nghĩa AI','Suy luận hợp lý',[]),
 (6,'Lịch sử','hệ chuyên gia',[]),
 (9,'Xử lý tiếng nói','Nhận dạng tiếng nói',[]),
 (10,'Siri','Apple',[]),
 (11,'ngôn ngữ tự nhiên','dịch máy',[]),
 (12,'chatbot','ngữ cảnh',[]),
 (13,'thị giác máy tính','Học sâu',[]),
 (18,'không gian trạng thái','không gian tìm kiếm',[]),
 (21,'Bài toán 8 số','3x3',[]),
 (23,'hai bình','4 lít',[]),
 (24,'trạng thái','(0, 0)',[]),
 (27,'tìm kiếm','blind searches',[]),
 (28,'Depth First Search','DFS',[18,27]),
 (30,'Breadth First Search','BFS',[18,27]),
 (37,'Phân rã bài toán','vấn đề con',[]),
 (48,'tốt nhất','hàm đánh giá',[30,45]),
 (51,'leo đồi','hàm đánh giá',[28,45]),
 (55,'tri thức','Cú pháp + Ngữ nghĩa + Cơ chế suy diễn',[]),
]

def norm(value):
    return unicodedata.normalize('NFC', str(value)).casefold()

def build_golden(source):
    deck = Presentation(source)
    cases = []
    for index, (slide, concept, evidence, prerequisites) in enumerate(ANNOTATIONS, 1):
        text = '\n'.join(s.text for s in deck.slides[slide-1].shapes if s.has_text_frame)
        assert norm(evidence) in norm(text), (slide, evidence)
        cases.append({'id': f'ai_{index:02d}', 'input_file': source.name,
            'input_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
            'source_text': text, 'expected_source': {'slide': slide},
            'expected_concept_contains': concept, 'expected_evidence': evidence,
            'expected_prerequisite_slides': prerequisites,
            'relationship_annotation': 'Curriculum expectation, not explicit slide edge',
            'expected_quiz': {'concept': concept, 'source_slide': slide, 'required_evidence': evidence},
            'expected_refusal': False, 'expected_mastery': [1.0,0.25],
            'expected_recommendation_slide': prerequisites[0] if prerequisites else slide})
    path = ROOT / 'eval/golden_set.json'
    backup = path.with_name('golden_set_legacy.json')
    if not backup.exists():
        backup.write_bytes(path.read_bytes())
    path.write_text(json.dumps(cases,ensure_ascii=False,indent=2),encoding='utf-8')

def run():
    parser = argparse.ArgumentParser()
    parser.add_argument('--build-golden', action='store_true')
    parser.add_argument('--llm', action='store_true', help='Call OpenAI for actual quiz generation (requires OPENAI_API_KEY)')
    args = parser.parse_args()
    source = ROOT / 'eval/slide AI.pptx'
    if args.build_golden:
        build_golden(source)
    cases = json.loads((ROOT / 'eval/golden_set.json').read_text(encoding='utf-8'))
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if any(c['input_sha256'] != digest for c in cases):
        raise ValueError('Sample changed; review golden annotations before rebuilding')
    original_db, original_upload = api.db_storage, api.UPLOAD_DIR
    results = []
    with tempfile.TemporaryDirectory(dir=ROOT / 'eval',prefix='slide_eval_') as directory:
        api.db_storage = StorageManager(str(Path(directory) / 'storage.json'))
        api.UPLOAD_DIR = directory
        try:
            with TestClient(app) as client:
                upload = client.post('/adaptive/upload',files={'file':(source.name,source.read_bytes(),
                    'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
                upload.raise_for_status()
                lid = upload.json()['lesson']['lesson_id']
                graph_response = client.get(f'/adaptive/knowledge-graph/{lid}')
                graph_response.raise_for_status()
                graph = graph_response.json()
                nodes = {n['slide']:n for n in graph['nodes']}
                for case in cases:
                    checks = {}
                    def record(stage,status,expected,actual,reason=''):
                        checks[stage] = dict(status=status,expected=expected,actual=actual,reason=reason)
                    def check(stage,passed,expected,actual):
                        record(stage,'PASS' if passed else 'FAIL',expected,actual)
                    slide = case['expected_source']['slide']
                    try:
                        check('upload',upload.status_code == 200,200,upload.status_code)
                        response = client.get(f'/adaptive/lessons/{lid}/slides/{slide}')
                        response.raise_for_status()
                        content = response.json()
                        check('citation_correctness',norm(case['expected_evidence']) in norm(content.get('text','')),
                              case['expected_evidence'],content)
                        node = nodes.get(slide)
                        check('concept_extraction_correctness',bool(node) and norm(case['expected_concept_contains']) in norm(node['label']),
                              case['expected_concept_contains'],node)
                        expected_edges = case['expected_prerequisite_slides']
                        if expected_edges:
                            check('relationship_correctness',all(any(e['from'] == nodes.get(p,{}).get('id') and
                                e['to'] == (node or {}).get('id') for e in graph['edges']) for p in expected_edges),expected_edges,graph['edges'])
                        else:
                            record('relationship_correctness','N/A',None,None,'No required edge annotated')
                        for stage in ['quiz_generation','groundedness','refusal_correctness']:
                            record(stage,'BLOCKED',case['expected_quiz'] if stage != 'refusal_correctness' else False,None,
                                   'No quiz generation/retrieval or refusal API implemented')
                        record('video_sync','N/A',None,None,'PPTX has no associated video/timestamps')
                        record('browser_click','BLOCKED',case['expected_recommendation_slide'],None,'Browser/FE interaction not implemented in this suite')
                        if not node:
                            raise ValueError(f'No graph node for slide {slide}')
                        cid = node['id']
                        if args.llm:
                            generated = client.post(f'/adaptive/lessons/{lid}/quizzes/generate',
                                json={'concept_id':cid,'count':1})
                            check('quiz_generation',generated.status_code == 200 and bool(generated.json().get('quizzes')),
                                  'one real LLM quiz',generated.json())
                            if generated.status_code == 200:
                                batch = generated.json()
                                check('refusal_correctness',batch['refused'] == case['expected_refusal'],
                                      case['expected_refusal'],batch['refused'])
                                if batch['quizzes']:
                                    quiz = batch['quizzes'][0]
                                    check('evidence_verbatim',quiz['evidence'] in content['text'],
                                          'verbatim source evidence',quiz['evidence'])
                                    record('groundedness','BLOCKED','semantic answer supported by evidence',None,
                                           'Verbatim evidence verified; semantic judge/reviewer still required')
                                    wrong = (quiz['correct_option'] + 1) % 4
                                    submission = client.post('/adaptive/quiz/answer',json={
                                        'quiz_id':quiz['quiz_id'],'student_id':case['id']+'_llm','selected_option':wrong})
                                    check('llm_quiz_wrong_answer',submission.status_code == 200 and
                                          submission.json().get('is_correct') is False and submission.json().get('new_mastery') == 0,
                                          {'is_correct':False,'new_mastery':0},submission.json())
                        mapping = client.get(f'/adaptive/concepts/by-slide/{slide}',params={'lesson_id':lid}).json()
                        check('slide_sync',cid in [c['concept_id'] for c in mapping['concepts']],cid,mapping)
                        scores = []
                        for index,(correct,weight) in enumerate([(True,1),(False,3)]):
                            answer = client.post('/adaptive/quiz/answer',params={'student_id':case['id']},json={
                                'quiz_id':f"fixture_{case['id']}_{index}",'concept_id':cid,'is_correct':correct,'difficulty':weight})
                            answer.raise_for_status()
                            scores.append(answer.json()['new_mastery'])
                        check('mastery_update',scores == case['expected_mastery'],case['expected_mastery'],scores)
                        summary = client.get(f"/adaptive/mastery/{case['id']}").json()
                        state = summary['mastery_summary'][cid]['state']
                        check('weak_concept',state == 'weak','weak',state)
                        response = client.get(f"/adaptive/recommendation/{case['id']}",params={'current_concept_id':cid})
                        response.raise_for_status()
                        recommendation = response.json()
                        destination = recommendation['source'].get('slide')
                        check('adaptive_recommendation_correctness',destination == case['expected_recommendation_slide'],
                              case['expected_recommendation_slide'],recommendation)
                        target = client.get(f'/adaptive/lessons/{lid}/slides/{destination}')
                        check('recommendation_destination_api',target.status_code == 200 and target.json().get('found') is True,
                              'existing slide content',target.json())
                    except Exception as exc:
                        record('execution','ERROR','complete API path',str(exc))
                    status = 'FAIL' if any(c['status'] in {'FAIL','ERROR'} for c in checks.values()) else 'BLOCKED'
                    results.append({'case_id':case['id'],'slide':slide,'e2e_status':status,'checks':checks})
                    print(f"{case['id']} slide {slide}: {status}")
                    for stage,value in checks.items():
                        if value['status'] in {'FAIL','ERROR','BLOCKED'}:
                            print(f"  {stage}: {value['status']} | expected={value['expected']} | actual={value['actual']} | {value['reason']}")
        finally:
            api.db_storage,api.UPLOAD_DIR = original_db,original_upload
    metrics = {}
    for result in results:
        for stage,value in result['checks'].items():
            counts = metrics.setdefault(stage,{s:0 for s in ['PASS','FAIL','BLOCKED','ERROR','N/A']})
            counts[value['status']] += 1
    for counts in metrics.values():
        executed = counts['PASS'] + counts['FAIL'] + counts['ERROR']
        counts['executed_pass_rate'] = round(100*counts['PASS']/executed,2) if executed else None
    report = {'mode':'real_pptx_api_evaluation','input_file':str(source),'sha256':digest,
        'total_cases':len(cases),'e2e_complete_passes':0,'e2e_completion_rate':0,
        'limitations':'Attempts are fixtures, not backend-generated quizzes. No expected edges injected. Video N/A. Missing capabilities BLOCKED.',
        'metrics':metrics,'cases':results}
    (ROOT / 'eval/backend_eval_results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(metrics,ensure_ascii=False,indent=2))
    return 1  # incomplete E2E must fail CI

if __name__ == '__main__':
    sys.exit(run())
