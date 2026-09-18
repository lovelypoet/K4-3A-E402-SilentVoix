import hashlib
import os
import re
import uuid
from typing import List, Dict, Any, Optional, Tuple
from urllib.parse import urlparse, parse_qs
from .models import Chunk

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from youtube_transcript_api import YouTubeTranscriptApi
except ImportError:
    YouTubeTranscriptApi = None

try:
    from faster_whisper import WhisperModel
except ImportError:
    WhisperModel = None

import urllib.request

# ---------------------------------------------------------------------------
# Ingestion status codes
# ---------------------------------------------------------------------------
INGESTION_OK            = "INGESTION_OK"
NO_EXTRACTABLE_TEXT     = "NO_EXTRACTABLE_TEXT"
OCR_REQUIRED            = "OCR_REQUIRED"
TRANSCRIPTION_REQUIRED  = "TRANSCRIPTION_REQUIRED"
EMPTY_SOURCE            = "EMPTY_SOURCE"
INGESTION_FAILED        = "INGESTION_FAILED"
DUPLICATE_SOURCE        = "DUPLICATE_SOURCE"

LOCAL_VIDEO_EXTENSIONS  = {".mp4", ".mov", ".avi", ".webm", ".mkv", ".mp3", ".wav", ".m4a"}



class DocumentIngestor:
    """
    Handles REAL ingestion and text extraction from actual PDF, PPTX, Transcript, YouTube URLs,
    strictly preserving Slide numbers and Timestamps.
    """

    @staticmethod
    def _chunk_id(prefix: str, lesson_id: str, idx: int) -> str:
        safe_lesson = re.sub(r"[^\w\-]+", "_", lesson_id)
        return f"{prefix}_{safe_lesson}_{idx:03d}"

    # ------------------------------------------------------------------
    # Source identity helpers
    # ------------------------------------------------------------------

    @staticmethod
    def extract_youtube_video_id(url: str) -> Optional[str]:
        """Extracts 11-character video ID from various YouTube URL formats."""
        if not url:
            return None
        pattern = r"(?:v=|\/|embed\/|shorts\/)([0-9A-Za-z_-]{11})"
        match = re.search(pattern, url)
        if match:
            return match.group(1)
        # bare youtu.be/ID
        parsed = urlparse(url)
        if "youtu.be" in (parsed.netloc or ""):
            candidate = (parsed.path or "").lstrip("/").split("/")[0]
            if len(candidate) == 11:
                return candidate
        return None

    @classmethod
    def youtube_source_key(cls, url: str) -> Optional[str]:
        """Return canonical source_key for a YouTube URL, ignoring timestamps."""
        vid = cls.extract_youtube_video_id(url)
        if vid:
            return f"youtube:{vid}"
        return None

    @staticmethod
    def file_sha256(file_path: str) -> Optional[str]:
        """Return SHA-256 hex digest of a local file's contents."""
        try:
            h = hashlib.sha256()
            with open(file_path, "rb") as f:
                for block in iter(lambda: f.read(65536), b""):
                    h.update(block)
            return h.hexdigest()
        except Exception:
            return None

    @classmethod
    def file_source_key(cls, file_path: str) -> Optional[str]:
        """Return a content-based source_key for an uploaded file."""
        digest = cls.file_sha256(file_path)
        if digest:
            return f"file:{digest}"
        return None

    @staticmethod
    def pdf_is_likely_scanned(file_path: str, min_chars_per_page: int = 20) -> bool:
        """
        Return True if pypdf extracts < min_chars_per_page on average
        — likely a scanned/image-only PDF.
        """
        if not pypdf:
            return False
        try:
            reader = pypdf.PdfReader(file_path)
            if not reader.pages:
                return True
            total_chars = sum(len((p.extract_text() or "").strip()) for p in reader.pages)
            avg = total_chars / len(reader.pages)
            return avg < min_chars_per_page
        except Exception:
            return True

    @staticmethod
    def validate_chunks(chunks: List[Chunk]) -> Tuple[bool, str]:
        """
        Post-extraction validation.
        Returns (ok: bool, error_code: str).
        A result with 0 chunks is NOT silently accepted.
        """
        if not chunks:
            return False, NO_EXTRACTABLE_TEXT
        meaningful = [c for c in chunks if len(c.text.strip()) >= 10]
        if not meaningful:
            return False, NO_EXTRACTABLE_TEXT
        return True, INGESTION_OK

    @staticmethod
    def stamp_chunks(
        chunks: List[Chunk],
        *,
        document_id: str,
        source_key: Optional[str] = None,
        source_group: Optional[str] = None,
        source_type: Optional[str] = None,
    ) -> List[Chunk]:
        """Apply identity fields after the document/lesson IDs are allocated."""
        for chunk in chunks:
            chunk.document_id = document_id
            chunk.source_key = source_key
            chunk.source_group = source_group or source_key
            chunk.source_type = source_type
        return chunks

    @classmethod
    def ingest_local_video_file(
        cls,
        file_path: str,
        lesson_id: str = "lesson_01",
        *,
        document_id: Optional[str] = None,
        source_key: Optional[str] = None,
        model_size: Optional[str] = None,
        device: Optional[str] = None,
        compute_type: Optional[str] = None,
    ) -> List[Chunk]:
        """Transcribe a local video/audio file into timestamped chunks.

        faster-whisper is optional because model download/runtime support is
        environment-dependent. Callers should report TRANSCRIPTION_REQUIRED
        when it is not installed, rather than registering an empty READY lesson.
        """
        if not WhisperModel:
            raise RuntimeError(f"{TRANSCRIPTION_REQUIRED}: faster-whisper is not installed")
        if not os.path.isfile(file_path):
            raise FileNotFoundError(file_path)

        model = WhisperModel(
            model_size or os.getenv("WHISPER_MODEL", "small"),
            device=device or os.getenv("WHISPER_DEVICE", "auto"),
            compute_type=compute_type or os.getenv("WHISPER_COMPUTE_TYPE", "int8"),
        )
        segments, _info = model.transcribe(file_path, vad_filter=True)
        chunks: List[Chunk] = []
        for idx, segment in enumerate(segments, start=1):
            text = (segment.text or "").strip()
            if not text:
                continue
            chunks.append(Chunk(
                chunk_id=cls._chunk_id("chunk_video", lesson_id, idx),
                lesson_id=lesson_id,
                document_id=document_id,
                text=text,
                start_time=round(float(segment.start), 2),
                end_time=round(float(segment.end), 2),
                source_key=source_key,
                source_group=source_key,
                source_type="video",
            ))
        valid, status = cls.validate_chunks(chunks)
        if not valid:
            raise RuntimeError(status)
        return chunks


    @classmethod
    def ingest_youtube_url(cls, video_url: str, lesson_id: str = "lesson_01") -> List[Chunk]:
        """Fetches REAL transcript and timestamps directly from a YouTube video URL."""
        if not YouTubeTranscriptApi:
            print("[DocumentIngestor] youtube_transcript_api not installed.")
            return []

        video_id = cls.extract_youtube_video_id(video_url)
        if not video_id:
            print(f"[DocumentIngestor] Could not extract video ID from {video_url}")
            return []

        chunks = []
        try:
            raw_transcript = YouTubeTranscriptApi().fetch(video_id, languages=["vi", "en"])

            current_start = None
            current_end = 0.0
            current_text_parts = []
            chunk_idx = 1

            for item in raw_transcript:
                text = getattr(item, "text", "").strip()
                start = float(getattr(item, "start", 0.0))
                duration = float(getattr(item, "duration", 0.0))
                end = start + duration

                if not text:
                    continue

                if current_start is None:
                    current_start = start

                current_text_parts.append(text)
                current_end = end

                if (end - current_start) >= 30.0:
                    combined_text = " ".join(current_text_parts).strip()
                    chunks.append(Chunk(
                        chunk_id=cls._chunk_id("chunk_yt", lesson_id, chunk_idx),
                        lesson_id=lesson_id,
                        text=combined_text,
                        start_time=round(current_start, 2),
                        end_time=round(current_end, 2),
                    ))
                    chunk_idx += 1
                    current_start = None
                    current_text_parts = []

            if current_text_parts and current_start is not None:
                combined_text = " ".join(current_text_parts).strip()
                chunks.append(Chunk(
                    chunk_id=cls._chunk_id("chunk_yt", lesson_id, chunk_idx),
                    lesson_id=lesson_id,
                    text=combined_text,
                    start_time=round(current_start, 2),
                    end_time=round(current_end, 2),
                ))

            print(
                f"[DocumentIngestor] Grouped {len(raw_transcript)} subtitle items into "
                f"{len(chunks)} 30-second chunks for YouTube video {video_id}"
            )
        except Exception as e:
            print(f"[DocumentIngestor] Failed to fetch transcript for YouTube video {video_id}: {e}")

        return chunks

    @classmethod
    def ingest_remote_document(
        cls,
        url: str,
        lesson_id: str = "lesson_01",
        upload_dir: Optional[str] = None,
    ) -> List[Chunk]:
        """
        Download remote PDF/PPTX/DOCX hoặc Google Docs/Slides/Drive (public) rồi ingest.
        """
        if not url:
            return []

        download_url = cls._resolve_download_url(url)
        if not download_url:
            raise ValueError(
                "Unsupported remote URL. Cần .pdf/.pptx/.docx trực tiếp hoặc Google Docs/Slides/Drive."
            )

        if upload_dir is None:
            upload_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads")
        os.makedirs(upload_dir, exist_ok=True)

        # Prefer format hint from export URL
        ext = ".pdf"
        dl_lower = download_url.lower()
        path_ext = os.path.splitext((urlparse(download_url).path or "").lower())[1]
        if path_ext in {".pdf", ".pptx", ".docx"}:
            ext = path_ext
        elif "format=pptx" in dl_lower or dl_lower.endswith("/export/pptx"):
            ext = ".pptx"
        elif "format=txt" in dl_lower:
            ext = ".txt"
        elif "format=docx" in dl_lower:
            ext = ".docx"
        elif "format=pdf" in dl_lower or "/export/pdf" in dl_lower:
            ext = ".pdf"

        filename = f"remote_{lesson_id}_{uuid.uuid4().hex[:8]}{ext}"
        dest = os.path.join(upload_dir, filename)

        req = urllib.request.Request(
            download_url,
            headers={"User-Agent": "SilentVoix-LessonStudio/1.0"},
        )
        with urllib.request.urlopen(req, timeout=45) as resp:
            data = resp.read()
            content_type = (resp.headers.get("Content-Type") or "").lower()
            # Google đôi khi trả HTML login thay vì file
            if "text/html" in content_type or (data[:200].lstrip().lower().startswith(b"<!doctype") or data[:100].lstrip().lower().startswith(b"<html")):
                raise ValueError(
                    "Google trả về trang HTML (file private hoặc cần đăng nhập). "
                    "Hãy bật 'Anyone with the link can view'."
                )
            if "presentationml" in content_type or "powerpoint" in content_type:
                ext = ".pptx"
                dest = dest.rsplit(".", 1)[0] + ".pptx"
            elif "wordprocessingml" in content_type or "officedocument.wordprocessing" in content_type:
                ext = ".docx"
                dest = dest.rsplit(".", 1)[0] + ".docx"
            elif "pdf" in content_type:
                ext = ".pdf"
                dest = dest.rsplit(".", 1)[0] + ".pdf"
            elif "text/plain" in content_type:
                ext = ".txt"
                dest = dest.rsplit(".", 1)[0] + ".txt"

        if not data or len(data) < 50:
            raise ValueError("Downloaded file is empty or too small.")

        with open(dest, "wb") as f:
            f.write(data)

        if ext == ".pdf":
            return cls.ingest_pdf_file(dest, lesson_id=lesson_id)
        if ext == ".pptx":
            return cls.ingest_pptx_file(dest, lesson_id=lesson_id)
        if ext == ".docx":
            return cls.ingest_docx_file(dest, lesson_id=lesson_id)
        if ext == ".txt":
            text = data.decode("utf-8", errors="ignore")
            return cls.process_raw_text(text, lesson_id=lesson_id)
        raise ValueError(f"Unsupported downloaded file type: {ext}")

    @staticmethod
    def _resolve_download_url(url: str) -> Optional[str]:
        """Convert Google Docs/Slides/Drive share links to export/download URLs."""
        parsed = urlparse(url)
        path = parsed.path or ""
        lower = url.lower()

        # Direct file links
        if path.lower().endswith((".pdf", ".pptx", ".docx")):
            return url

        # Google Docs: /document/d/<ID>/...
        m = re.search(r"/document/d/([a-zA-Z0-9-_]+)", url)
        if m:
            file_id = m.group(1)
            return f"https://docs.google.com/document/d/{file_id}/export?format=pdf"

        # Google Slides: /presentation/d/<ID>/...
        m = re.search(r"/presentation/d/([a-zA-Z0-9-_]+)", url)
        if m:
            file_id = m.group(1)
            return f"https://docs.google.com/presentation/d/{file_id}/export/pdf"

        # Google Drive file: /file/d/<ID>/...
        m = re.search(r"/file/d/([a-zA-Z0-9-_]+)", url)
        if m:
            file_id = m.group(1)
            return f"https://drive.google.com/uc?export=download&id={file_id}"

        # open?id= / uc?id=
        qs = parse_qs(parsed.query or "")
        if "id" in qs and ("drive.google.com" in lower or "docs.google.com" in lower):
            file_id = qs["id"][0]
            return f"https://drive.google.com/uc?export=download&id={file_id}"

        return None

    @staticmethod
    def extract_clean_concepts_from_chunks(chunks: List[Chunk]) -> Dict[str, Dict[str, Any]]:
        """
        Extracts clean, high-level Knowledge Concepts from parsed slide or transcript chunks.
        Concept IDs are lesson-scoped to avoid collisions across lessons.
        """
        extracted_concepts = {}
        extracted_mappings = {}
        lesson_id = chunks[0].lesson_id if chunks else "lesson_01"
        safe_lesson = re.sub(r"[^\w\-]+", "_", lesson_id)

        topic_keywords = [
            ("ai agent", "Khái niệm AI Agent", "Lập trình và kiến trúc hệ thống AI Agent"),
            ("prompt", "Prompt Engineering", "Kỹ thuật tối ưu hóa prompt cho LLM"),
            ("tool", "LLM Tool & Function Calling", "Tích hợp công cụ và hàm cho mô hình ngôn ngữ"),
            ("gradient descent", "Thuật toán Gradient Descent", "Tối ưu hóa hàm mất mát bằng Gradient Descent"),
            ("loss function", "Hàm Mất Mát (Loss Function)", "Đánh giá sai số của mô hình học máy"),
            ("linear regression", "Hồi Quy Tuyến Tính (Linear Regression)", "Mô hình dự đoán tuyến tính cơ bản"),
            ("supervised", "Học Có Giám Sát (Supervised Learning)", "Huấn luyện mô hình với dữ liệu gán nhãn"),
            ("unsupervised", "Học Không Giám Sát (Unsupervised Learning)", "Phân cụm và tìm cấu trúc dữ liệu"),
            ("python", "Lập Trình Python Cho AI", "Ứng dụng ngôn ngữ Python trong phát triển AI"),
            ("machine learning", "Nền Tảng Machine Learning", "Tổng quan về Học Máy và các ứng dụng"),
        ]

        for c in chunks:
            if c.slide:
                lines = [line.strip() for line in c.text.split("\n") if line.strip()]
                if lines:
                    title_line = lines[0]
                    clean_title = re.sub(
                        r"^(?:slide\s*\d+:?|\d+[\.\)]\s*|[-*•]\s*)",
                        "",
                        title_line,
                        flags=re.IGNORECASE,
                    ).strip()
                    if clean_title and len(clean_title) >= 3:
                        if len(clean_title) > 50:
                            clean_title = clean_title[:47] + "..."
                        cid = f"c_{safe_lesson}_slide_{c.slide}"
                        extracted_concepts[cid] = {
                            "concept_id": cid,
                            "name": clean_title,
                            "description": c.text[:200],
                            "lesson_id": lesson_id,
                        }
                        extracted_mappings[cid] = {
                            "concept_id": cid,
                            "lesson_id": lesson_id,
                            "slide": c.slide,
                            "start_time": None,
                            "end_time": None,
                        }

            elif c.text:
                lower_text = c.text.lower()
                for kw, topic_name, desc in topic_keywords:
                    if kw in lower_text:
                        cid = f"c_{safe_lesson}_{kw.replace(' ', '_')}"
                        if cid not in extracted_concepts:
                            extracted_concepts[cid] = {
                                "concept_id": cid,
                                "name": topic_name,
                                "description": desc,
                                "lesson_id": lesson_id,
                            }
                            extracted_mappings[cid] = {
                                "concept_id": cid,
                                "lesson_id": lesson_id,
                                "slide": None,
                                "start_time": c.start_time,
                                "end_time": c.end_time,
                            }

        return {"concepts": extracted_concepts, "source_mappings": extracted_mappings}

    @classmethod
    def ingest_pdf_file(cls, file_path: str, lesson_id: str = "lesson_01") -> List[Chunk]:
        """
        Parses PDF page-by-page; page == chunk.page == slide for source mapping.

        Pipeline:
            pypdf extraction
              ↓
            meaningful text? YES → return chunks
                             NO  → raise ValueError(OCR_REQUIRED)
        """
        if not pypdf:
            raise ImportError("pypdf is required to parse PDF files. Please install pypdf.")

        reader = pypdf.PdfReader(file_path)
        chunks = []
        total_text_chars = 0
        for idx, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            total_text_chars += len(text)
            if text:
                chunks.append(Chunk(
                    chunk_id=cls._chunk_id("chunk_pdf", lesson_id, idx + 1),
                    lesson_id=lesson_id,
                    text=text,
                    slide=idx + 1,
                    page=idx + 1,
                ))

        # If no text was extracted at all, classify the failure
        if not chunks:
            n_pages = len(reader.pages)
            if n_pages == 0:
                raise ValueError(f"{EMPTY_SOURCE}: PDF has no pages.")
            raise ValueError(
                f"{OCR_REQUIRED}: PDF '{os.path.basename(file_path)}' ({n_pages} pages) "
                "produced 0 text characters — likely scanned or image-only. "
                "Use an OCR tool and re-upload as text-based PDF."
            )

        return chunks


    @classmethod
    def ingest_pptx_file(cls, file_path: str, lesson_id: str = "lesson_01") -> List[Chunk]:
        """Parses PPTX slide-by-slide; slide_number = slide_idx + 1."""
        if not Presentation:
            raise ImportError("python-pptx is required to parse PPTX files. Please install python-pptx.")

        chunks = []
        prs = Presentation(file_path)
        for idx, slide in enumerate(prs.slides):
            slide_text_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_text_parts.append(shape.text.strip())

            full_slide_text = "\n".join(slide_text_parts).strip()
            if full_slide_text:
                chunks.append(Chunk(
                    chunk_id=cls._chunk_id("chunk_pptx", lesson_id, idx + 1),
                    lesson_id=lesson_id,
                    text=full_slide_text,
                    slide=idx + 1,
                ))
        return chunks

    @classmethod
    def ingest_docx_file(cls, file_path: str, lesson_id: str = "lesson_01") -> List[Chunk]:
        """Parses DOCX; mỗi heading/paragraph block → 1 'slide' chunk."""
        if not DocxDocument:
            raise ImportError("python-docx is required to parse DOCX files. Please install python-docx.")

        doc = DocxDocument(file_path)
        blocks: List[str] = []
        current: List[str] = []

        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text:
                if current:
                    blocks.append("\n".join(current))
                    current = []
                continue
            style_name = (para.style.name if para.style else "") or ""
            is_heading = style_name.lower().startswith("heading")
            if is_heading and current:
                blocks.append("\n".join(current))
                current = [text]
            else:
                current.append(text)

        if current:
            blocks.append("\n".join(current))

        # Fallback: nếu không tách được block, gộp toàn bộ
        if not blocks:
            all_text = "\n".join((p.text or "").strip() for p in doc.paragraphs if (p.text or "").strip())
            if all_text:
                blocks = [all_text]

        chunks = []
        for idx, block in enumerate(blocks):
            chunks.append(Chunk(
                chunk_id=cls._chunk_id("chunk_docx", lesson_id, idx + 1),
                lesson_id=lesson_id,
                text=block,
                slide=idx + 1,
            ))
        return chunks

    @staticmethod
    def ingest_slides_data(slides_data: List[Dict[str, Any]], lesson_id: str = "lesson_01") -> List[Chunk]:
        """Convert structured slide records into source-preserving chunks."""
        chunks = []
        for index, item in enumerate(slides_data or [], start=1):
            text = str(item.get("text") or "").strip()
            if not text:
                continue
            slide_number = item.get("slide", index)
            chunks.append(Chunk(
                chunk_id=f"chunk_slide_{index:03d}",
                lesson_id=lesson_id,
                text=text,
                slide=int(slide_number) if slide_number is not None else None,
            ))
        return chunks

    @staticmethod
    def ingest_transcript_json(transcript_data: List[Dict[str, Any]], lesson_id: str = "lesson_01") -> List[Chunk]:
        """
        Ingests a transcript list with start_time, end_time, and text.
        """
        chunks = []
        for i, item in enumerate(transcript_data):
            text = item.get("text", "").strip()
            if not text:
                continue
            chunk = Chunk(
                chunk_id=f"chunk_tr_{i+1:03d}",
                lesson_id=lesson_id,
                text=text,
                start_time=float(item.get("start_time", 0)),
                end_time=float(item.get("end_time", 0))
            )
            chunks.append(chunk)
        return chunks

    @staticmethod
    def ingest_srt_vtt_file(file_path: str, lesson_id: str = "lesson_01") -> List[Chunk]:
        """
        Parses real subtitle files (.srt or .vtt) to extract text and timestamps.
        """
        chunks = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()

        def timestamp_to_seconds(ts_str: str) -> float:
            ts_str = ts_str.replace(',', '.')
            parts = ts_str.split(':')
            if len(parts) == 3:
                h, m, s = parts
                return float(h) * 3600 + float(m) * 60 + float(s)
            elif len(parts) == 2:
                m, s = parts
                return float(m) * 60 + float(s)
            return 0.0

        pattern = re.compile(r'(\d{1,2}:\d{2}:\d{2}[\.,]\d{1,3}|\d{2}:\d{2}[\.,]\d{1,3})\s*-->\s*(\d{1,2}:\d{2}:\d{2}[\.,]\d{1,3}|\d{2}:\d{2}[\.,]\d{1,3})')
        
        lines = content.splitlines()
        idx = 0
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            match = pattern.search(line)
            if match:
                start_sec = timestamp_to_seconds(match.group(1))
                end_sec = timestamp_to_seconds(match.group(2))
                text_lines = []
                i += 1
                while i < len(lines) and lines[i].strip() and not pattern.search(lines[i]):
                    text_lines.append(lines[i].strip())
                    i += 1
                text = " ".join(text_lines).strip()
                if text:
                    idx += 1
                    chunks.append(Chunk(
                        chunk_id=f"chunk_sub_{idx:03d}",
                        lesson_id=lesson_id,
                        text=text,
                        start_time=start_sec,
                        end_time=end_sec
                    ))
            else:
                i += 1
        return chunks

    @staticmethod
    def process_raw_text(text: str, lesson_id: str = "lesson_01", default_slide: int = 1) -> List[Chunk]:
        """
        Fallback chunker for plain text.
        """
        if not text or not text.strip():
            return []
        
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        chunks = []
        for idx, p in enumerate(paragraphs):
            chunks.append(Chunk(
                chunk_id=f"chunk_txt_{idx+1:03d}",
                lesson_id=lesson_id,
                text=p,
                slide=default_slide
            ))
        return chunks
